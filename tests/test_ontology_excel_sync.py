# tests/test_ontology_excel_sync.py
"""
온톨로지 엑셀 ↔ SoT 정합 (2026-10-01).

엑셀(handoff/ontology_v4.9/spec/CCOP_Ontology_V4.9_node_edge_attrs.xlsx)의 카탈로그·속성 사전 시트는
scripts/build_ontology_excel.py 가 SoT + 속성 사전 파일에서 생성한다. 손으로 고치거나 SoT 만 바꾸고
재생성을 빼먹으면 여기서 실패한다 — `python3 scripts/build_ontology_excel.py` 로 다시 만들면 된다.
"""
import importlib.util
import os

import openpyxl
import pytest

from app.services.ontology_service import KICSCrimeDomainOntology as O

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
_spec = importlib.util.spec_from_file_location('build_ontology_excel', os.path.join(ROOT, 'scripts', 'build_ontology_excel.py'))
G = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(G)


@pytest.fixture(scope='module')
def wb():
    return openpyxl.load_workbook(G.XLSX, read_only=True)


def _rows(ws):
    return [['' if v is None else v for v in r] for r in ws.iter_rows(min_row=2, values_only=True) if any(v is not None for v in r)]


@pytest.mark.parametrize('prefix,title,header,rows', G.generated_sheets(), ids=lambda x: x if isinstance(x, str) else '')
def test_generated_sheet_matches_sot(wb, prefix, title, header, rows):
    assert title in wb.sheetnames, f'시트 없음/제목 불일치: {title} — scripts/build_ontology_excel.py 재실행'
    got = _rows(wb[title])
    want = [['' if v is None else v for v in r] for r in rows]
    assert got == want, f'{title} 이 SoT 와 다름 — scripts/build_ontology_excel.py 재실행'


def test_catalog_covers_sot(wb):
    edges = {r[1] for r in _rows(wb[f'엣지카탈로그({len(O.RELATIONSHIPS)}종)'])}
    assert edges == set(O.RELATIONSHIPS)
    nodes = {r[2] for r in _rows(wb[f'노드카탈로그({len(O.ENTITIES)}종)'])}
    assert nodes == {v['label'] for v in O.ENTITIES.values()}


def test_every_attribute_has_description():
    need = [r[:2] for _, _, _, rows in G.generated_sheets() for r in rows if G.NEED in r]
    assert not need, f'설명 없는 속성 — app/services/ontology_attribute_dictionary.json 에 추가: {need}'


def test_candidates_are_not_in_sot():
    """후보(미확정) 속성은 SoT 정의에 없어야 한다 — 등재되면 후보 목록에서 빼야 함."""
    node = {v['label']: set(v.get('properties', [])) | set(v.get('attributes', [])) for v in O.ENTITIES.values()}
    for c in O.ATTRIBUTE_DICTIONARY.get('candidates', []):
        if c['kind'] == 'node':
            assert c['attr'] not in node[c['target']], c
        else:
            assert c['attr'] not in (O.RELATIONSHIPS[c['target']].get('properties') or []), c


def test_event_node_identifier_consistent():
    """NODE_ID_STANDARD 정경 식별자 = ENTITIES properties (이벤트 노드 event_id 모순 재발 방지)."""
    for v in O.ENTITIES.values():
        canon = O.NODE_ID_STANDARD.get(v['label'], {}).get('canonical_field', '')
        if canon and not canon.startswith('('):
            assert canon in v.get('properties', []), (v['label'], canon, v.get('properties'))
