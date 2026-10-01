#!/usr/bin/env python3
"""온톨로지 엑셀(노드·엣지 속성 정보)의 카탈로그·속성 사전 시트를 SoT 에서 생성한다.

SoT = app/services/ontology_service.py (무엇이 노드·엣지·속성인가)
    + app/services/ontology_attribute_dictionary.json (사람이 읽는 설명·분류·후보 속성)

생성 시트 (매번 통째로 다시 씀):
  엣지카탈로그(N종) · 노드카탈로그(N종) · 노드 속성 사전 · 엣지 속성 사전 · 후보 속성(미확정)
수기 시트 (건드리지 않음):
  변경이력 · 값 도메인 사전 · 탐색·적재 규칙 · 파생속성 등록부 · 공통 메타·안내

엑셀을 손으로 고치지 말고 SoT·사전 파일을 고친 뒤 이 스크립트를 돌린다 (2026-10-01 — 엑셀이 SoT 와
따로 관리되며 sameAs·폐지 속성 등이 남던 문제의 구조적 해결). tests/test_ontology_excel_sync.py 가 검사.

실행: python3 scripts/build_ontology_excel.py [--xlsx 경로]
"""
import argparse
import copy
import os
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
from app.services.ontology_service import KICSCrimeDomainOntology as O  # noqa: E402

XLSX = os.path.join(ROOT, 'handoff', 'ontology_v4.9', 'spec', 'CCOP_Ontology_V4.9_node_edge_attrs.xlsx')
D = O.ATTRIBUTE_DICTIONARY
COMMON_NODE = set(D.get('common_attrs', {}))            # 노드 공통 메타 — 노드별 행에서 제외(공통 메타 시트)
NODE_META_EDGE = {'source_id', 'rec_created'}           # 카탈로그 '주요 속성'에서 생략
NEED = '(설명 필요)'
LABEL = {k: v['label'] for k, v in O.ENTITIES.items()}


def _labels(expr):
    out = []
    for t in str(expr).split('|'):
        t = t.strip()
        out.append('(any)' if t == 'Any' else LABEL.get(t, t))
    return ' | '.join(out)


def _derived_by_node():
    m = {}
    for a, v in O.DERIVED_PROPERTY_REGISTRY.items():
        for l in str(v.get('node', '')).split('|'):
            m.setdefault(l.strip(), []).append(a)
    return m


def edge_catalog_rows():
    order = list(D.get('edge_catalog', {})) + [e for e in O.RELATIONSHIPS if e not in D.get('edge_catalog', {})]
    rows = []
    for e in order:
        if e not in O.RELATIONSHIPS:
            continue
        v = O.RELATIONSHIPS[e]
        group, note = D.get('edge_catalog', {}).get(e, ['(미분류)', ''])
        props = [p for p in (v.get('properties') or []) if p not in NODE_META_EDGE]
        rows.append([group, e, _labels(v.get('domain')), _labels(v.get('range')),
                     ', '.join(props) or '-', v.get('meaning') or '', note or ''])
    return rows


def node_catalog_rows():
    rows = []
    for i, v in enumerate(O.ENTITIES.values(), start=1):
        l = v['label']
        nid = O.NODE_ID_STANDARD.get(l, {})
        core = D.get('node_core', {}).get(l) or ', '.join(a for a in v.get('attributes', []) if a not in COMMON_NODE)[:80]
        rows.append([i, v.get('layer'), l, nid.get('canonical_field') or ', '.join(v.get('properties', [])),
                     nid.get('default_format') or '', core])
    return rows


def node_attr_rows():
    desc, derived = D.get('node_attrs', {}), _derived_by_node()
    rows = []
    for v in O.ENTITIES.values():
        l = v['label']
        attrs = [a for a in v.get('properties', []) + v.get('attributes', []) if a not in COMMON_NODE]
        attrs += [a for a in derived.get(l, []) if a not in attrs]
        for a in attrs:
            t, d, vt = (desc.get(l, {}).get(a) or [None, NEED, None])
            if a in derived.get(l, []) and not str(d).startswith('[파생]'):
                d = f'[파생] {d}'
            rows.append([l, a, t or '', d or NEED, vt or ''])
    return rows


def edge_attr_rows():
    used = {}
    for e, v in O.RELATIONSHIPS.items():
        for p in v.get('properties') or []:
            used.setdefault(p, []).append(e)
    desc, common = D.get('edge_attrs', {}), D.get('common_attrs', {})
    order = [a for a in desc if a in used or a in O.EDGE_META_SCHEMA] + [a for a in used if a not in desc]
    rows = []
    for a in order:
        t, d, vt = desc.get(a) or common.get(a) or [None, NEED, None]
        es = used.get(a, [])
        if a in O.EDGE_META_SCHEMA:
            where = '(엣지 공통 메타)'                      # EDGE_META_SCHEMA — 전 엣지에 붙을 수 있음
        else:
            where = ', '.join(es[:8]) + (f' 외 {len(es) - 8}' if len(es) > 8 else '')
        rows.append([a, t or '', d or NEED, where, vt or ''])
    return rows


def candidate_rows():
    return [['노드' if c['kind'] == 'node' else '엣지', c['target'], c['attr'], c.get('type') or '',
             c.get('desc') or '', c.get('value_type') or '', c.get('basis') or '']
            for c in D.get('candidates', [])]


HEADERS = {
    'edge_catalog': ['분류', '엣지', '노드1(출발)', '노드2(도착)', '주요 속성', '의미', '비고'],
    'node_catalog': ['#', '레이어', '라벨', '표준 식별자', 'id_format', '핵심 속성'],
    'node_attrs': ['노드', '속성', '타입', '설명', '값유형'],
    'edge_attrs': ['속성', '타입', '설명', '주 사용 엣지', '값유형'],
    'candidates': ['구분', '대상', '속성', '타입', '설명', '값유형', '판정 근거'],
}


def generated_sheets():
    """(시트 접두어, 새 제목, 헤더, 행) — 테스트도 이 함수로 기대값을 만든다."""
    return [
        ('엣지카탈로그', f'엣지카탈로그({len(O.RELATIONSHIPS)}종)', HEADERS['edge_catalog'], edge_catalog_rows()),
        ('노드카탈로그', f'노드카탈로그({len(O.ENTITIES)}종)', HEADERS['node_catalog'], node_catalog_rows()),
        ('노드 속성 사전', '노드 속성 사전', HEADERS['node_attrs'], node_attr_rows()),
        ('엣지 속성 사전', '엣지 속성 사전', HEADERS['edge_attrs'], edge_attr_rows()),
        ('후보 속성', '후보 속성(미확정)', HEADERS['candidates'], candidate_rows()),
    ]


def _rewrite(ws, title, header, rows, data_style):
    ws.title = title
    ws.delete_rows(1, ws.max_row)
    for j, h in enumerate(header, start=1):
        c = ws.cell(1, j, h)
        if data_style.get('head', {}).get(j):
            c._style = copy.copy(data_style['head'][j])
    for i, r in enumerate(rows, start=2):
        for j, x in enumerate(r, start=1):
            c = ws.cell(i, j, x)
            if data_style.get('body', {}).get(j):
                c._style = copy.copy(data_style['body'][j])


def build(path=XLSX):
    import openpyxl
    wb = openpyxl.load_workbook(path)
    for prefix, title, header, rows in generated_sheets():
        ws = next((w for w in wb.worksheets if w.title.startswith(prefix)), None)
        if ws is None:                       # 후보 속성 시트 신설 — 엣지 속성 사전 뒤, 스타일은 그 시트에서
            ref = next(w for w in wb.worksheets if w.title.startswith('엣지 속성 사전'))
            ws = wb.create_sheet(title, wb.worksheets.index(ref) + 1)
            src = ref
            for k, d in ref.column_dimensions.items():
                ws.column_dimensions[k].width = d.width
            ws.column_dimensions['G'].width = 48
        else:
            src = ws
        style = {'head': {c.column: c._style for c in src[1] if c.has_style},
                 'body': {c.column: c._style for c in src[2] if c.has_style} if src.max_row >= 2 else {}}
        _rewrite(ws, title, header, rows, style)
        ws.freeze_panes = 'A2'
    wb.save(path)
    return [(t, len(r)) for _, t, _, r in generated_sheets()]


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--xlsx', default=XLSX)
    out = build(ap.parse_args().xlsx)
    for t, n in out:
        print(f'  {t:20} {n}행')
    need = sum(1 for _, _, _, rows in generated_sheets() for r in rows if NEED in r)
    print(f'생성: {ap.parse_args().xlsx}' + (f'  ⚠ 설명 필요 {need}행' if need else ''))
