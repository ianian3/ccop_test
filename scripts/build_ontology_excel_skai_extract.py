#!/usr/bin/env python3
"""V4.9 추출조건판(SKAI) 엑셀 생성 — SKAI 9/17판(V4.8)을 바탕으로 V4.9 변경분만 반영하고 변경 위치를 표시한다.

원칙 (2026-10-01 사용자 지시):
  · 앞쪽 SKAI 작성 시트(추출가능_노드속성·추출가능_엣지속성·쿼리연산자)와 SKAI_추가근거는 손대지 않는다
  · 변경이력 · 엣지카탈로그 · 노드카탈로그 · 노드 속성 사전 · 엣지 속성 사전 및 그 뒤 시트(값 도메인 사전 ·
    탐색·적재 규칙 · 파생속성 등록부 · 공통 메타·안내)만 V4.9 정본 내용으로 바꾼다. V4.9 신설 시트
    (후보 속성(미확정) · 공통 속성 그룹)는 엣지 속성 사전 뒤에 넣는다
  · 바뀐 곳 표시: 변경 셀 노란색(옛 값은 'V4.9 변경' 열에), 신규 행 초록색, 삭제 행은 시트 끝에 빨간 취소선으로 남김

입력: ~/Downloads/CCOP_Ontology_V4.8_노드엣지속성정보_추출조건_SKAI_20260917.xlsx (저장소 밖 원본)
      handoff/ontology_v4.9/spec/CCOP_Ontology_V4.9_node_edge_attrs.xlsx (V4.9 정본 — scripts/build_ontology_excel.py 로 먼저 재생성)
출력: ~/Downloads/CCOP_Ontology_V4.9_노드엣지속성정보_추출조건_SKAI_20261001.xlsx
"""
import copy
import os
import sys

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
D = os.path.expanduser('~/Downloads/')
BASE = D + 'CCOP_Ontology_V4.8_노드엣지속성정보_추출조건_SKAI_20260917.xlsx'
V49 = os.path.join(ROOT, 'handoff', 'ontology_v4.9', 'spec', 'CCOP_Ontology_V4.9_node_edge_attrs.xlsx')
OUT = D + 'CCOP_Ontology_V4.9_노드엣지속성정보_추출조건_SKAI_20261001.xlsx'

KEEP = ('추출가능_노드속성', '추출가능_엣지속성', '쿼리연산자', 'SKAI_추가근거')
# (V4.8 시트 접두어, V4.9 시트 접두어, 행 키 함수)  — 키로 V4.8 행과 V4.9 행을 짝짓는다
SHEETS = [
    ('변경이력', '변경이력', lambda r: (r[0],)),
    ('엣지카탈로그', '엣지카탈로그', lambda r: (r[1],)),
    ('노드카탈로그', '노드카탈로그', lambda r: (r[2],)),
    ('노드 속성 사전', '노드 속성 사전', lambda r: (r[0], r[1])),
    ('엣지 속성 사전', '엣지 속성 사전', lambda r: (r[0],)),
    (None, '후보 속성', None),
    (None, '공통 속성 그룹', None),
    ('값 도메인 사전', '값 도메인 사전', lambda r: (r[0], r[2])),        # (속성, 값) — 적용 엣지 표기(sameAs→same_as) 변경은 '변경'으로
    ('탐색·적재 규칙', '탐색·적재 규칙', lambda r: (r[0],)),
    ('파생속성 등록부', '파생속성 등록부', lambda r: (r[0],)),
    ('공통 메타·안내', '공통 메타·안내', lambda r: (r[0],)),
]
FILL_CHG = PatternFill('solid', fgColor='FFF2CC')     # 변경 셀
FILL_NEW = PatternFill('solid', fgColor='E2EFDA')     # 신규 행
FILL_DEL = PatternFill('solid', fgColor='FCE4E4')     # 삭제 행
MARK_COL_TITLE = 'V4.9 변경'


def _sheet(wb, prefix):
    return next((w for w in wb.worksheets if w.title.startswith(prefix)), None)


def _norm(v):
    return '' if v is None or str(v).strip() in ('', 'None') else str(v).strip()


def _rows(ws):
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if any(_norm(v) for v in r):
            out.append(list(r))
    return out


def _keyed(rows, keyf, fill_down_first=False):
    """키 → 행 목록(같은 키 여러 행이면 순서대로). 값 도메인 사전처럼 첫 열이 이어지는 행은 위 값을 채운다."""
    out, last = {}, None
    for r in rows:
        r = list(r)
        if fill_down_first:
            if _norm(r[0]):
                last = r[0]
            else:
                r = [last] + r[1:]
        out.setdefault(tuple(_norm(x) for x in keyf(r)), []).append(r)
    return out


def build():
    wb = openpyxl.load_workbook(BASE)
    src = openpyxl.load_workbook(V49)
    summary = {}
    insert_after = None
    for old_pre, new_pre, keyf in SHEETS:
        new_ws = _sheet(src, new_pre)
        old_ws = _sheet(wb, old_pre) if old_pre else None
        if old_ws is None:                                   # V4.9 신설 시트 — 직전 시트 뒤에 만든다
            ref = _sheet(wb, insert_after)
            old_ws = wb.create_sheet(new_ws.title, wb.worksheets.index(ref) + 1)
            for k, d in new_ws.column_dimensions.items():
                old_ws.column_dimensions[k].width = d.width
            style_src = ref
        else:
            style_src = old_ws
        insert_after = new_pre
        head_style = {c.column: copy.copy(c._style) for c in style_src[1] if c.has_style}
        body_style = {c.column: copy.copy(c._style) for c in style_src[2] if c.has_style} if style_src.max_row >= 2 else {}
        last_col = max(head_style or [1])

        header = [c.value for c in new_ws[1]]
        new_rows = _rows(new_ws)
        old_rows = _rows(old_ws) if old_pre else []
        fd = new_pre == '값 도메인 사전'
        old_by = _keyed(old_rows, keyf, fd) if keyf else {}
        new_by = _keyed(new_rows, keyf, fd) if keyf else {}

        old_ws.title = new_ws.title
        old_ws.delete_rows(1, old_ws.max_row)
        ncol = len(header) + 1                               # + 'V4.9 변경' 열

        def put(i, j, v, style_from):
            c = old_ws.cell(i, j, v)
            st = style_from.get(j) or style_from.get(min(j, last_col))
            if st is not None:
                c._style = copy.copy(st)
            return c

        for j, h in enumerate(header + [MARK_COL_TITLE], start=1):
            put(1, j, h, head_style)
        old_ws.column_dimensions[openpyxl.utils.get_column_letter(ncol)].width = 30
        old_ws.cell(1, ncol).comment = Comment('노란 셀: 값 변경(옛 값은 이 열) · 초록 행: V4.9 신규 · 빨간 취소선 행: V4.9 삭제',
                                               'CCOP')
        n_new = n_chg = n_del = 0
        i = 2
        seen = {}
        for r in new_rows:
            rr = list(r)
            if fd and not _norm(rr[0]) and i > 2:            # 값 도메인 사전 이어지는 행 — 키는 위 값으로
                rr_key = None
            k = None
            if keyf:
                kr = list(rr)
                if fd and not _norm(kr[0]):
                    kr[0] = last_dom
                else:
                    last_dom = kr[0] if fd else None
                k = tuple(_norm(x) for x in keyf(kr))
            occ = seen.get(k, 0)
            seen[k] = occ + 1
            olds = old_by.get(k, []) if k is not None else []
            old = olds[occ] if occ < len(olds) else None
            for j in range(len(header)):
                put(i, j + 1, rr[j] if j < len(rr) else None, body_style)
            mark = put(i, ncol, None, body_style)
            if old is None:
                for j in range(1, ncol + 1):
                    old_ws.cell(i, j).fill = FILL_NEW
                mark.value = '신규'
                n_new += 1
            else:
                diffs = []
                for j in range(len(header)):
                    ov = old[j] if j < len(old) else None
                    nv = rr[j] if j < len(rr) else None
                    if j == 0 and fd and not _norm(nv):
                        continue
                    if _norm(ov) != _norm(nv):
                        old_ws.cell(i, j + 1).fill = FILL_CHG
                        diffs.append(f'{header[j]}: {_norm(ov) or "(빈칸)"}')
                if diffs:
                    mark.value = '변경 — 구 ' + ' / '.join(diffs)
                    mark.fill = FILL_CHG
                    n_chg += 1
            i += 1
        # 삭제 행 — V4.8 에만 있는 키
        for k, olds in old_by.items():
            extra = olds[len(new_by.get(k, [])):]
            for r in extra:
                for j in range(len(header)):
                    c = put(i, j + 1, r[j] if j < len(r) else None, body_style)
                    c.fill = FILL_DEL
                    c.font = Font(strike=True, color='9C0006')
                m = put(i, ncol, 'V4.9 삭제', body_style)
                m.fill = FILL_DEL
                m.font = Font(bold=True, color='9C0006')
                n_del += 1
                i += 1
        old_ws.freeze_panes = 'A2'
        summary[new_ws.title] = {'신규': n_new, '변경': n_chg, '삭제': n_del}

    # 시트 순서: 보존 시트(앞 3종) → V4.9 시트 → SKAI_추가근거
    order = [wb[t] for t in KEEP[:3]] + [_sheet(wb, p) for _, p, _ in SHEETS] + [wb[KEEP[3]]]
    wb._sheets = order
    wb.save(OUT)
    return summary


if __name__ == '__main__':
    for sheet, s in build().items():
        print(f'  {sheet:22} 신규 {s["신규"]:>3} · 변경 {s["변경"]:>3} · 삭제 {s["삭제"]:>3}')
    print('생성:', OUT)
